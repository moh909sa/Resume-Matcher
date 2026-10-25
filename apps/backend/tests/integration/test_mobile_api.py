"""Mobile APIs over real engine storage and a mocked external Supabase HTTP API."""

import copy
import json
from datetime import datetime, timezone
from io import BytesIO
from typing import Any

import httpx
import pytest
import respx
from docx import Document
from pydantic import SecretStr

from app.config import settings
from app.database import Database
from app.main import app
from app.resume_ownership import resume_owner
from app.routers import resumes

A = "760c63e9-c782-4631-a1fb-471e988740a1"
B = "760c63e9-c782-4631-a1fb-471e988740b2"
URL = "https://project.supabase.test"
KEY = "sb_secret_mobile_test_only"


class SupabaseHTTP:
    """Privileged API: filters must come from the app, not implicit mock RLS."""

    def __init__(self) -> None:
        self.rows: dict[str, dict[str, Any]] = {}
        self.objects: dict[str, bytes] = {}
        self.calls: list[httpx.Request] = []
        self.fail_rpc = False

    def handle(self, request: httpx.Request) -> httpx.Response:
        self.calls.append(request)
        assert request.headers["apikey"] == KEY
        path = request.url.path
        if path == "/auth/v1/user":
            token = request.headers.get("authorization")
            user = {"Bearer token-a": A, "Bearer token-b": B}.get(token)
            if user is None:
                return httpx.Response(401, json={"message": "Expired or invalid JWT"})
            return httpx.Response(200, json={"id": user, "is_anonymous": False})
        # Do not carry the user's Authorization header into privileged writes.
        assert "authorization" not in request.headers
        payload = (
            json.loads(request.content)
            if request.content
            and request.headers.get("content-type", "").startswith("application/json")
            else None
        )
        if path == "/rest/v1/rpc/reserve_user_resume":
            if (
                sum(
                    row["user_id"] == payload["p_user_id"] for row in self.rows.values()
                )
                >= 5
            ):
                return httpx.Response(409, json={"message": "quota"})
            row = {
                "id": payload["p_resume_id"],
                "engine_resume_id": payload["p_resume_id"],
                "user_id": payload["p_user_id"],
                "original_filename": payload["p_filename"],
                "storage_path": payload["p_storage_path"],
                "storage_bucket": "resume-files",
                "mime_type": payload["p_mime_type"],
                "file_size_bytes": payload["p_file_size"],
                "processing_status": "pending",
                "is_default": False,
            }
            if row["id"] in self.rows:
                return httpx.Response(409, json={"message": "duplicate"})
            now = datetime.now(timezone.utc).isoformat()
            row.update(created_at=now, updated_at=now)
            self.rows[row["id"]] = row
            return httpx.Response(201, json=[row])
        if path == "/rest/v1/user_resumes":
            rows = list(self.rows.values())
            for key in ("user_id", "id"):
                value = request.url.params.get(key)
                if value:
                    assert value.startswith("eq.")
                    rows = [row for row in rows if row[key] == value[3:]]
            if request.method == "GET":
                return httpx.Response(200, json=rows)
            if request.method == "PATCH":
                for row in rows:
                    row.update(payload)
                return httpx.Response(204)
            if request.method == "DELETE":
                for row in rows:
                    del self.rows[row["id"]]
                return httpx.Response(204)
        if path == "/rest/v1/rpc/sync_user_resume_registry":
            if self.fail_rpc:
                return httpx.Response(503, json={"message": f"internal {KEY}"})
            states = {row["resume_id"]: row for row in payload["p_snapshot"]}
            for row in self.rows.values():
                if row["user_id"] == payload["p_user_id"]:
                    state = states.get(row["engine_resume_id"])
                    row["is_default"] = bool(state and state["is_default_master"])
                    if state:
                        row["processing_status"] = state["processing_status"]
            return httpx.Response(204)
        prefix = "/storage/v1/object/resume-files/"
        if request.method == "POST" and path.startswith(prefix):
            key = path[len(prefix) :]
            assert request.headers["x-upsert"] == "false"
            self.objects[key] = request.content
            return httpx.Response(200, json={"Key": key})
        prefix = "/storage/v1/object/authenticated/resume-files/"
        if request.method == "GET" and path.startswith(prefix):
            key = path[len(prefix) :]
            if key not in self.objects:
                return httpx.Response(404)
            return httpx.Response(200, content=self.objects[key])
        if request.method == "DELETE" and path == "/storage/v1/object/resume-files":
            for key in payload["prefixes"]:
                self.objects.pop(key, None)
            return httpx.Response(200, json=[])
        raise AssertionError(f"Unmocked Supabase operation: {request.method} {path}")


@pytest.fixture
def upstream(
    monkeypatch: pytest.MonkeyPatch, respx_mock: respx.MockRouter
) -> SupabaseHTTP:
    monkeypatch.setattr(settings, "mobile_api_enabled", True)
    monkeypatch.setattr(settings, "supabase_url", URL)
    monkeypatch.setattr(settings, "supabase_secret_key", SecretStr(KEY))
    server = SupabaseHTTP()
    respx_mock.route(url__startswith=URL).mock(side_effect=server.handle)
    return server


@pytest.fixture
async def client(upstream: SupabaseHTTP) -> httpx.AsyncClient:
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as value:
        yield value


@pytest.fixture
def original() -> bytes:
    doc = Document()
    doc.add_heading("Alex Morgan", 0)
    doc.add_paragraph(
        "Software engineer with experience building Python APIs and cloud services."
    )
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


@pytest.fixture
def parsed(monkeypatch: pytest.MonkeyPatch, sample_resume: dict[str, Any]) -> None:
    async def parse(_markdown: str) -> dict[str, Any]:
        return copy.deepcopy(sample_resume)

    monkeypatch.setattr(resumes, "parse_resume_to_json", parse)


async def upload(
    client: httpx.AsyncClient, original: bytes, token: str = "token-a"
) -> httpx.Response:
    return await client.post(
        "/api/v1/mobile/resumes/upload",
        headers={"Authorization": f"Bearer {token}"},
        files={
            "file": (
                "original.docx",
                original,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )


@pytest.mark.integration
async def test_upload_retains_original_and_reuses_processing(
    client: httpx.AsyncClient,
    upstream: SupabaseHTTP,
    original: bytes,
    parsed: None,
    isolated_db: Database,
) -> None:
    response = await upload(client, original)
    assert response.status_code == 201, response.text
    row = response.json()
    assert row["processing_status"] == "ready"
    assert row["processed_resume"]["personalInfo"]["name"] == "Jane Doe"
    assert row["id"] == row["engine_resume_id"]
    stored = upstream.rows[row["id"]]
    assert stored["user_id"] == A
    assert upstream.objects[stored["storage_path"]] == original
    assert stored["storage_path"].startswith(f"{A}/{row['id']}/")
    with resume_owner(A):
        assert (await isolated_db.get_resume(row["id"]))["processing_status"] == "ready"
    assert await isolated_db.get_resume(row["id"]) is None
    assert KEY not in response.text and "token-a" not in response.text


@pytest.mark.integration
async def test_lists_and_all_id_operations_are_isolated(
    client: httpx.AsyncClient,
    upstream: SupabaseHTTP,
    original: bytes,
    parsed: None,
) -> None:
    a = (await upload(client, original, "token-a")).json()
    b = (await upload(client, original, "token-b")).json()
    for token, expected in (("token-a", a), ("token-b", b)):
        headers = {"Authorization": f"Bearer {token}"}
        listed = await client.get("/api/v1/mobile/resumes", headers=headers)
        assert listed.status_code == 200
        assert [row["id"] for row in listed.json()] == [expected["id"]]
    headers = {"Authorization": "Bearer token-b"}
    for method, suffix, body in (
        ("GET", "", None),
        ("GET", "/original", None),
        ("DELETE", "", None),
        ("POST", "/retry-processing", None),
        ("POST", "/default", None),
        ("PATCH", "", {"display_name": "Stolen"}),
    ):
        response = await client.request(
            method,
            f"/api/v1/mobile/resumes/{a['id']}{suffix}",
            headers=headers,
            json=body,
        )
        assert response.status_code == 404, (method, suffix, response.text)
    assert a["id"] in upstream.rows and b["id"] in upstream.rows
    downloaded = await client.get(
        f"/api/v1/mobile/resumes/{a['id']}/original",
        headers={"Authorization": "Bearer token-a"},
    )
    assert downloaded.content == original
    assert "attachment" in downloaded.headers["content-disposition"]


@pytest.mark.integration
async def test_parse_failure_keeps_file_and_retry_uses_same_id(
    client: httpx.AsyncClient,
    upstream: SupabaseHTTP,
    original: bytes,
    monkeypatch: pytest.MonkeyPatch,
    sample_resume: dict[str, Any],
) -> None:
    async def fail(_markdown: str) -> dict[str, Any]:
        raise RuntimeError("Provider unavailable")

    monkeypatch.setattr(resumes, "parse_resume_to_json", fail)
    response = await upload(client, original)
    assert response.status_code == 201
    row = response.json()
    assert row["processing_status"] == "failed"
    assert len(upstream.objects) == 1

    async def success(_markdown: str) -> dict[str, Any]:
        return sample_resume

    monkeypatch.setattr(resumes, "parse_resume_to_json", success)
    response = await client.post(
        f"/api/v1/mobile/resumes/{row['id']}/retry-processing",
        headers={"Authorization": "Bearer token-a"},
    )
    assert response.status_code == 200, response.text
    assert response.json()["id"] == row["id"]
    assert response.json()["processing_status"] == "ready"
    assert len(upstream.rows) == len(upstream.objects) == 1


@pytest.mark.integration
async def test_missing_engine_is_rebuilt_from_private_original(
    client: httpx.AsyncClient,
    upstream: SupabaseHTTP,
    original: bytes,
    parsed: None,
    isolated_db: Database,
) -> None:
    row = (await upload(client, original)).json()
    with resume_owner(A):
        await isolated_db.delete_resume(row["id"])
    headers = {"Authorization": "Bearer token-a"}
    response = await client.get(f"/api/v1/mobile/resumes/{row['id']}", headers=headers)
    assert response.json()["recovery_required"] is True
    response = await client.post(
        f"/api/v1/mobile/resumes/{row['id']}/retry-processing", headers=headers
    )
    assert response.status_code == 200, response.text
    assert (
        response.json()["id"] == row["id"]
        and response.json()["processing_status"] == "ready"
    )


@pytest.mark.integration
async def test_delete_removes_bytes_registry_and_engine_only_for_owner(
    client: httpx.AsyncClient,
    upstream: SupabaseHTTP,
    original: bytes,
    parsed: None,
    isolated_db: Database,
) -> None:
    row = (await upload(client, original)).json()
    response = await client.delete(
        f"/api/v1/mobile/resumes/{row['id']}",
        headers={"Authorization": "Bearer token-a"},
    )
    assert response.status_code == 204
    assert not upstream.rows and not upstream.objects
    with resume_owner(A):
        assert await isolated_db.get_resume(row["id"]) is None


@pytest.mark.integration
async def test_upstream_failure_does_not_leak_secrets_and_upload_can_be_recovered(
    client: httpx.AsyncClient,
    upstream: SupabaseHTTP,
    original: bytes,
    parsed: None,
) -> None:
    upstream.fail_rpc = True
    response = await upload(client, original)
    assert response.status_code == 503
    assert KEY not in response.text
    resume_id = response.json()["resume_id"]
    assert resume_id in upstream.rows and len(upstream.objects) == 1
    upstream.fail_rpc = False
    response = await client.get(
        f"/api/v1/mobile/resumes/{resume_id}",
        headers={"Authorization": "Bearer token-a"},
    )
    assert (
        response.status_code == 200 and response.json()["processing_status"] == "ready"
    )


@pytest.mark.integration
@pytest.mark.parametrize("token", [None, "expired", "forged-other-project"])
async def test_missing_and_invalid_sessions_are_rejected(
    client: httpx.AsyncClient,
    upstream: SupabaseHTTP,
    token: str | None,
) -> None:
    headers = {"Authorization": f"Bearer {token}"} if token else {}
    response = await client.get("/api/v1/mobile/resumes", headers=headers)
    assert response.status_code == 401
    assert not upstream.rows and not upstream.objects


@pytest.mark.integration
@pytest.mark.parametrize(
    "path,method",
    [
        ("/api/v1/resumes/list?include_master=true", "GET"),
        ("/api/v1/config/reset", "POST"),
        ("/api/v1/config/api-keys", "GET"),
        ("/api/v1/jobs/upload", "POST"),
        ("/auth/callback", "GET"),
    ],
)
async def test_mobile_mode_closes_single_user_apis(
    client: httpx.AsyncClient, path: str, method: str
) -> None:
    response = await client.request(
        method, path, headers={"Authorization": "Bearer token-a"}
    )
    assert response.status_code == 404


@pytest.mark.integration
async def test_missing_secret_fails_closed(
    client: httpx.AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(settings, "supabase_secret_key", SecretStr(""))
    assert (
        await client.get(
            "/api/v1/mobile/resumes", headers={"Authorization": "Bearer token-a"}
        )
    ).status_code == 503
    assert (await client.get("/api/v1/resumes/list")).status_code == 404
    assert (await client.get("/api/v1/health")).status_code == 200


@pytest.mark.integration
async def test_oversized_upload_never_reaches_storage(
    client: httpx.AsyncClient, upstream: SupabaseHTTP
) -> None:
    response = await client.post(
        "/api/v1/mobile/resumes/upload",
        headers={"Authorization": "Bearer token-a"},
        files={"file": ("large.pdf", b"x" * (4 * 1024 * 1024 + 1), "application/pdf")},
    )
    assert response.status_code == 413
    assert not upstream.rows and not upstream.objects


@pytest.mark.integration
async def test_client_cannot_choose_another_owner_or_engine_id(
    client: httpx.AsyncClient, upstream: SupabaseHTTP, original: bytes, parsed: None
) -> None:
    response = await client.post(
        "/api/v1/mobile/resumes/upload",
        headers={"Authorization": "Bearer token-a"},
        data={"user_id": B, "engine_resume_id": B},
        files={
            "file": (
                "cv.docx",
                original,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert response.status_code == 201, response.text
    row = response.json()
    assert upstream.rows[row["id"]]["user_id"] == A and row["engine_resume_id"] != B


@pytest.mark.integration
async def test_quota_is_per_account_and_does_not_store_a_sixth_file(
    client: httpx.AsyncClient,
    upstream: SupabaseHTTP,
    original: bytes,
    parsed: None,
) -> None:
    for _ in range(5):
        assert (await upload(client, original)).status_code == 201
    response = await upload(client, original)
    assert response.status_code == 409
    assert len(upstream.rows) == len(upstream.objects) == 5
    assert (await upload(client, original, "token-b")).status_code == 201


@pytest.mark.integration
async def test_rename_default_and_delete_promote_only_within_account(
    client: httpx.AsyncClient,
    upstream: SupabaseHTTP,
    original: bytes,
    parsed: None,
) -> None:
    first = (await upload(client, original)).json()
    second = (await upload(client, original)).json()
    other = (await upload(client, original, "token-b")).json()
    headers = {"Authorization": "Bearer token-a"}
    response = await client.patch(
        f"/api/v1/mobile/resumes/{second['id']}",
        headers=headers,
        json={"display_name": "My updated CV"},
    )
    assert (
        response.status_code == 200
        and response.json()["display_name"] == "My updated CV"
    )
    response = await client.post(
        f"/api/v1/mobile/resumes/{second['id']}/default", headers=headers
    )
    assert response.status_code == 200 and response.json()["is_default"]
    assert not upstream.rows[first["id"]]["is_default"]
    assert upstream.rows[other["id"]]["is_default"]
    assert (
        await client.delete(f"/api/v1/mobile/resumes/{second['id']}", headers=headers)
    ).status_code == 204
    assert upstream.rows[first["id"]]["is_default"]
    assert upstream.rows[other["id"]]["is_default"]
